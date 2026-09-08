import csv
import io
from typing import List, Tuple

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models.models import DeviceTypeEnum
from app.schemas.common import DeviceCreate
from app.services.device import create_device
from app.services.device_sheet import HEADER_FIELDS


def parse_upload(file: UploadFile) -> List[dict]:
    filename = (file.filename or "").lower()
    content = file.file.read()
    if filename.endswith(".csv"):
        text = content.decode("utf-8-sig")
        return list(csv.DictReader(io.StringIO(text)))
    if filename.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(content))
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            wb.close()
            return []
        wb.close()
        headers = [str(x).strip() if x else "" for x in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:] if any(row)]
    raise HTTPException(status_code=400, detail="仅支持CSV或XLSX文件")



def normalize_row(row: dict) -> DeviceCreate:
    row = {HEADER_FIELDS.get(str(key).strip(), str(key).strip()): (value if value is not None else "")
           for key, value in row.items()}
    for field in ("asset_number", "name", "device_type", "model"):
        if not str(row.get(field, "")).strip():
            raise HTTPException(status_code=400, detail=f"必填项缺失: {field}")
    ports = []
    if row.get("ports"):
        for port_name in str(row.get("ports")).split("|"):
            port_name = port_name.strip()
            if port_name:
                ports.append({"name": port_name})
    admin_ids = []
    if row.get("administrator_ids"):
        admin_ids = [int(x) for x in str(row.get("administrator_ids")).split("|") if str(x).strip()]
    try:
        return DeviceCreate(
            asset_number=str(row.get("asset_number", "")).strip(),
            name=str(row.get("name", "")).strip(),
            device_type=DeviceTypeEnum(str(row.get("device_type", "server")).strip() or "server"),
            model=str(row.get("model", "")).strip(),
            serial_number=str(row.get("serial_number", "")).strip() or None,
            purpose=str(row.get("purpose", "")).strip() or None,
            status=str(row.get("status", "planned")).strip() or "planned",
            ip_address=str(row.get("ip_address", "")).strip() or None,
            rack_id=int(row.get("rack_id")) if row.get("rack_id") not in (None, "") else None,
            start_u=int(row.get("start_u")) if row.get("start_u") not in (None, "") else None,
            u_height=int(row.get("u_height") or 1),
            cpu=str(row.get("cpu", "")).strip() or None,
            memory_gb=int(row.get("memory_gb")) if row.get("memory_gb") not in (None, "") else None,
            gpu=str(row.get("gpu", "")).strip() or None,
            storage_desc=str(row.get("storage_desc", "")).strip() or None,
            operating_system=str(row.get("operating_system", "")).strip() or None,
            notes=str(row.get("notes", "")).strip() or None,
            administrator_ids=admin_ids,
            ports=ports,
        )
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"导入行格式错误: {row}") from exc



def import_devices(db: Session, rows: List[dict], user) -> Tuple[int, List[str]]:
    imported = 0
    errors = []
    for index, row in enumerate(rows, start=2):
        try:
            payload = normalize_row(row)
            create_device(db, payload, user)
            imported += 1
        except Exception as exc:  # noqa: BLE001
            db.rollback()
            detail = getattr(exc, "detail", str(exc))
            errors.append(f"第{index}行: {detail}")
    return imported, errors
