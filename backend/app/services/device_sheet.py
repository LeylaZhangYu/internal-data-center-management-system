"""Shared column definitions for device exports, templates and imports."""
import csv
import io

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [
    ("asset_number", "资产编号"), ("name", "设备名称"),
    ("device_type", "设备类型"), ("model", "型号"),
    ("serial_number", "序列号"), ("purpose", "用途"),
    ("status", "状态"), ("ip_address", "IP地址"),
    ("rack_id", "机柜ID"), ("start_u", "起始U位"),
    ("u_height", "占用U数"), ("cpu", "CPU"),
    ("memory_gb", "内存GB"), ("gpu", "GPU"),
    ("storage_desc", "硬盘"), ("operating_system", "操作系统"),
    ("notes", "备注"),
]
HEADER_FIELDS = {label: field for field, label in COLUMNS}
HEADER_FIELDS.update({"管理员ID": "administrator_ids", "端口": "ports"})


def sheet_response(file_format, devices=(), template=False):
    headers = [label for _, label in COLUMNS]
    rows = []
    for device in devices:
        values = [getattr(device, field) for field, _ in COLUMNS]
        rows.append([getattr(value, "value", value) if value is not None else "" for value in values])
    filename = "devices_template" if template else "devices"
    if file_format == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(headers)
        # Treat user-controlled text as text when opened in spreadsheet software.
        for row in rows:
            writer.writerow(["'" + v if isinstance(v, str) and v.startswith(("=", "+", "-", "@")) else v for v in row])
        content = output.getvalue().encode("utf-8-sig")
        media_type = "text/csv; charset=utf-8"
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "设备导入模板" if template else "设备清单"
        ws.append(headers)
        for row in rows:
            ws.append(row)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for index, cell in enumerate(ws[1], 1):
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="244768")
            ws.column_dimensions[get_column_letter(index)].width = 22
        ws.row_dimensions[1].height = 26
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        if template:
            help_sheet = wb.create_sheet("填写说明")
            for line in [
                "请在第一个工作表从第2行开始填写，不要修改表头。模板不包含示例设备。",
                "必填：资产编号、设备名称、设备类型、型号。资产编号不能与已有设备重复。",
                "设备类型：server / switch / storage / pdu / firewall / other。",
                "状态：planned / active / standby / maintenance / retired / off_shelf；留空为planned。",
                "机柜ID填写系统内已有机柜的数字ID，不是机柜编号。未上架时机柜ID和起始U位留空。",
                "上架时填写机柜ID、起始U位和占用U数；占用U数默认1，必须是正整数。",
                "内存GB填写整数；其余可选内容留空即可。设备类型和状态仍使用英文代码。",
                "导入用于新增设备，不会覆盖已有设备；导出数据重新导入前请检查资产编号。",
            ]:
                help_sheet.append([line])
            help_sheet.column_dimensions["A"].width = 115
        output = io.BytesIO()
        wb.save(output)
        wb.close()
        content = output.getvalue()
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return StreamingResponse(iter([content]), media_type=media_type,
                             headers={"Content-Disposition": f"attachment; filename={filename}.{file_format}"})
