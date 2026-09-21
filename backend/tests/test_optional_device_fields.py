import csv
import io

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import MetaData, create_engine, event, inspect, text
from sqlalchemy.orm import Session
from sqlalchemy.exc import IntegrityError

from app.core.database import Base
from app.core.schema_updates import ensure_optional_device_fields
from app.models.models import Area, DataCenter, Device, Rack
from app.schemas.common import DeviceCreate

FIELDS = ("asset_number", "model", "ip_address")


@pytest.mark.parametrize("empty", [None, "", "   "])
def test_create_multiple_empty_devices(client, admin_headers, empty):
    ids = []
    for index in range(2):
        payload = dict(name=f"可选字段设备{index}", device_type="server", rack_id=1, start_u=index + 1,
                       serial_number="", **dict.fromkeys(FIELDS, empty))
        response = client.post("/api/devices", json=payload, headers=admin_headers)
        assert response.status_code == 200, response.text
        assert all(response.json()[field] is None for field in FIELDS)
        ids.append(response.json()["id"])
    assert ids[0] != ids[1]
    assert client.get("/api/devices", headers=admin_headers).status_code == 200
    detail = client.get(f"/api/devices/{ids[0]}", headers=admin_headers)
    assert detail.status_code == 200, detail.text
    layout = client.get("/api/visualization/layout", headers=admin_headers)
    assert layout.status_code == 200, layout.text
    devices = layout.json()[0]["areas"][0]["racks"][0]["devices"]
    assert len(devices) == 2
    assert devices[0]["asset_number"] is None


def test_omitted_fields_and_clear_existing_values(client, admin_headers):
    bare = dict(name="未编号设备", device_type="server")
    created = client.post("/api/devices", json=bare, headers=admin_headers)
    assert created.status_code == 200, created.text
    assert all(created.json()[field] is None for field in FIELDS)
    populated = dict(bare, name="有编号设备", asset_number="TEST-001", model="R750", ip_address="10.90.0.1")
    saved = client.post("/api/devices", json=populated, headers=admin_headers)
    assert saved.status_code == 200, saved.text
    for field in ("asset_number", "ip_address"):
        duplicate = dict(bare, **{field: populated[field]})
        rejected = client.post("/api/devices", json=duplicate, headers=admin_headers)
        assert rejected.status_code == 400, rejected.text
    cleared = client.put(f'/api/devices/{saved.json()["id"]}',
                         json=dict(bare, **dict.fromkeys(FIELDS, "")), headers=admin_headers)
    assert cleared.status_code == 200, cleared.text
    assert all(cleared.json()[field] is None for field in FIELDS)
    # Cleared unique values can subsequently be used by another device.
    assert client.post("/api/devices", json=populated, headers=admin_headers).status_code == 200


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
@pytest.mark.parametrize("chinese", [True, False])
def test_import_export_optional_fields(client, admin_headers, extension, chinese):
    headers = ["设备名称", "设备类型", "资产编号", "型号", "IP地址"] if chinese else ["name", "device_type", *FIELDS]
    rows = [headers, ["导入设备1", "server", "", "", ""], ["导入设备2", "server", "", "", ""]]
    if extension == "csv":
        out = io.StringIO()
        csv.writer(out).writerows(rows)
        content = out.getvalue().encode("utf-8-sig")
    else:
        wb = Workbook()
        for row in rows:
            wb.active.append(row)
        out = io.BytesIO()
        wb.save(out)
        wb.close()
        content = out.getvalue()
    response = client.post("/api/devices/import", headers=admin_headers,
                           files={"file": (f"devices.{extension}", content)})
    assert response.status_code == 200, response.text
    assert response.json()["message"] == "导入完成，成功 2 条"
    exported = client.get(f"/api/devices/export/{extension}", headers=admin_headers)
    assert exported.status_code == 200, exported.text
    if extension == "csv":
        data = list(csv.DictReader(io.StringIO(exported.content.decode("utf-8-sig"))))
        assert len(data) == 2
        assert all(row[field] == "" for row in data for field in ("资产编号", "型号", "IP地址"))
    else:
        wb = load_workbook(io.BytesIO(exported.content))
        values = list(wb.active.values)
        assert len(values) == 3
        assert all(row[values[0].index(field)] is None for row in values[1:] for field in ("资产编号", "型号", "IP地址"))
        wb.close()
    template = client.get("/api/devices/template/xlsx", headers=admin_headers)
    assert template.status_code == 200
    wb = load_workbook(io.BytesIO(template.content))
    assert "资产编号、型号、IP地址可留空" in "".join(row[0] for row in wb["填写说明"].values)
    wb.close()


def test_required_schema_fields_remain_required():
    fields = DeviceCreate.__fields__
    assert fields["name"].required
    assert fields["device_type"].required
    assert all(not fields[field].required and fields[field].allow_none for field in FIELDS)


def test_existing_sqlite_upgrade_preserves_records_and_relationships(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'old.db'}", future=True)

    @event.listens_for(engine, "connect")
    def enable_foreign_keys(connection, _):
        connection.execute("PRAGMA foreign_keys=ON")

    old_metadata = MetaData()
    for table in Base.metadata.sorted_tables:
        table.to_metadata(old_metadata)
    old_metadata.tables["devices"].c.asset_number.nullable = False
    old_metadata.tables["devices"].c.model.nullable = False
    old_metadata.create_all(engine)
    with Session(engine) as db:
        dc = DataCenter(name="旧机房", location="旧位置", owner="管理员")
        db.add(dc)
        db.flush()
        area = Area(data_center_id=dc.id, name="旧区域")
        db.add(area)
        db.flush()
        rack = Rack(data_center_id=dc.id, area_id=area.id, code="OLD-RACK")
        db.add(rack)
        db.flush()
        device = Device(name="原有设备", asset_number="KEEP-001", model="旧型号", device_type="server", rack_id=rack.id)
        db.add(device)
        db.commit()
    with engine.begin() as connection:
        connection.execute(text("INSERT INTO device_placement_history(device_id,action,to_rack_id,u_height,created_at) VALUES (1,'mounted',1,1,CURRENT_TIMESTAMP)"))
        connection.execute(text("CREATE INDEX custom_device_name ON devices(name)"))
    ensure_optional_device_fields(engine)
    ensure_optional_device_fields(engine)
    assert all(col["nullable"] for col in inspect(engine).get_columns("devices") if col["name"] in FIELDS)
    assert "custom_device_name" in {index["name"] for index in inspect(engine).get_indexes("devices")}
    with engine.connect() as connection:
        assert connection.execute(text("SELECT asset_number, model FROM devices WHERE id=1")).one() == ("KEEP-001", "旧型号")
        assert connection.execute(text("SELECT device_id, to_rack_id FROM device_placement_history")).one() == (1, 1)
        assert not connection.exec_driver_sql("PRAGMA foreign_key_check").all()
        assert connection.exec_driver_sql("PRAGMA foreign_keys").scalar() == 1
    with Session(engine) as db:
        db.add_all([Device(name=f"空字段{i}", device_type="server") for i in range(2)])
        db.commit()
    with Session(engine) as db:
        db.add(Device(name="编号冲突", asset_number="KEEP-001", device_type="server"))
        with pytest.raises(IntegrityError):
            db.commit()
        db.rollback()
    assert len(list(tmp_path.glob("*.bak"))) == 1
    engine.dispose()
